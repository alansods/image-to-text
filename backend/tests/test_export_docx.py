"""Word gerado do documento conferido (spec document-export, design D10)."""

from pathlib import Path

import docx
import pytest
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app.export.docx_export import build_docx
from app.models.document import BBox
from tests.export_samples import sample_document


@pytest.fixture
def word_file(tmp_path: Path) -> docx.document.Document:
    document = sample_document(tmp_path)
    out = tmp_path / "out.docx"
    build_docx(document, tmp_path, out)
    return docx.Document(str(out))


def page_sections(doc: docx.document.Document) -> list:  # type: ignore[type-arg]
    """Seções que começam página nova (as contínuas só trocam colunas na mesma página)."""
    return [s for s in doc.sections if s.start_type != 0]  # 0 = CONTINUOUS


def all_text(doc: docx.document.Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [c.text for row in table.rows for c in row.cells]
    for s in doc.sections:
        parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return "\n".join(parts)


def paragraph(doc: docx.document.Document, text: str):  # type: ignore[no-untyped-def]
    return next(p for p in doc.paragraphs if p.text.strip() == text)


def test_one_page_per_original_page_with_same_sheet(word_file: docx.document.Document) -> None:
    sections = page_sections(word_file)
    assert len(sections) == 4
    assert [round(s.page_width.pt) for s in sections] == [595, 595, 595, 842]
    assert [round(s.page_height.pt) for s in sections] == [842, 842, 842, 595]
    assert sections[3].orientation == WD_ORIENT.LANDSCAPE


def test_title_centered_and_larger_paragraph_justified(word_file: docx.document.Document) -> None:
    title = paragraph(word_file, "CONTRATO DE SERVIÇOS")
    body = paragraph(word_file, "O CONTRATANTE declara")

    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert title.runs[0].font.size > body.runs[0].font.size


def test_bold_only_on_the_bold_word(word_file: docx.document.Document) -> None:
    body = paragraph(word_file, "O CONTRATANTE declara")
    bold = {run.text.strip() for run in body.runs if run.bold}
    assert bold == {"CONTRATANTE"}


def test_reviewed_text_is_exported(word_file: docx.document.Document) -> None:
    assert "assinado por SILVA" in all_text(word_file)


def test_list_keeps_markers_and_levels(word_file: docx.document.Document) -> None:
    first = paragraph(word_file, "a) recebimento")
    second = paragraph(word_file, "b) digitalização")
    assert second.paragraph_format.left_indent > (first.paragraph_format.left_indent or 0)


def test_table_is_a_real_table_with_merged_cell(word_file: docx.document.Document) -> None:
    [table] = word_file.tables
    assert (len(table.rows), len(table.columns)) == (4, 3)
    first = table.rows[0].cells[0]
    assert first.text == "Período"
    assert first._tc.tcPr.find(qn("w:gridSpan")).get(qn("w:val")) == "2"
    assert table.cell(3, 2).text == "3.200,00"


def test_every_figure_becomes_an_image(word_file: docx.document.Document) -> None:
    assert len(word_file.inline_shapes) == 3


def test_header_and_footer_are_repeated_page_parts(word_file: docx.document.Document) -> None:
    first = page_sections(word_file)[0]
    assert "Relatório mensal" in "\n".join(p.text for p in first.header.paragraphs)
    assert "Página 1" in "\n".join(p.text for p in first.footer.paragraphs)
    body = "\n".join(p.text for p in word_file.paragraphs)
    assert "Relatório mensal" not in body


def test_two_column_page_has_two_column_section(word_file: docx.document.Document) -> None:
    columns = [
        int(s._sectPr.find(qn("w:cols")).get(qn("w:num"), "1"))
        for s in word_file.sections
        if s._sectPr.find(qn("w:cols")) is not None
    ]
    assert 2 in columns
    texts = [p.text.strip() for p in word_file.paragraphs]
    assert texts.index("coluna direita texto") > texts.index("coluna esquerda texto")


def test_table_columns_keep_their_proportions(tmp_path: Path) -> None:
    document = sample_document(tmp_path)
    page = document.pages[0]
    table = next(b for b in page.blocks if b.cells)
    widths = [300, 1200, 300]  # coluna do meio 4x mais larga
    lefts = [300, 600, 1800]
    cells = [
        c.model_copy(
            update={
                "bbox": BBox(
                    x=lefts[c.col],
                    y=1600 + c.row * 100,
                    w=sum(widths[c.col : c.col + c.col_span]),
                    h=100,
                )
            }
        )
        for c in table.cells
    ]
    blocks = [b.model_copy(update={"cells": cells}) if b is table else b for b in page.blocks]
    document.pages[0] = page.model_copy(update={"blocks": blocks})
    out = tmp_path / "t.docx"
    build_docx(document, tmp_path, out)

    [word_table] = docx.Document(str(out)).tables
    row = word_table.rows[1].cells
    assert row[1].width > 3 * row[0].width
    grid = word_table._tbl.tblGrid.findall(qn("w:gridCol"))
    assert int(grid[1].get(qn("w:w"))) > 3 * int(grid[0].get(qn("w:w")))
