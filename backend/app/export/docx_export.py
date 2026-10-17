"""Word a partir do `Document` conferido (spec document-export, design D10).

Uma seção por página do original, com o mesmo tamanho de folha; cabeçalho e rodapé da página
viram cabeçalho e rodapé da seção; páginas com colunas usam seções contínuas com colunas;
cada palavra mantém negrito/itálico; listas mantêm o marcador original; tabelas com mescla;
figuras como imagem recortada do original, com largura proporcional.
"""

from collections.abc import Sequence
from pathlib import Path
from statistics import median
from typing import Any

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.models.document import ILLEGIBLE, Alignment, Block, BlockType, Document, Page, Word

FONT = "Times New Roman"  # Liberation Serif no servidor (mesmas medidas)
DEFAULT_SIZE = 11.0
MIN_SIZE, MAX_SIZE = 6.0, 40.0
# Margens da folha: vêm da posição do conteúdo no original, dentro destes limites.
MIN_MARGIN, MAX_MARGIN = 18.0, 90.0
# Bloco que ocupa ao menos isto da largura do conteúdo atravessa as colunas.
FULL_WIDTH = 0.6
LIST_INDENT = 18.0
# Entrelinha simples do Word já deixa ~0,3 do corpo entre linhas; o resto do vão vira espaço.
LEADING_GAP = 0.3
# Maior espaço vertical antes de um bloco (fração da altura da folha).
MAX_GAP = 0.6
# Cabeçalho/rodapé mais estreito que isto (fração da página) segue a posição na página.
SHORT_EDGE = 0.3
IMAGE_TYPES = {BlockType.FIGURE, BlockType.FORMULA}

_ALIGN = {
    Alignment.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
    Alignment.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    Alignment.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
    Alignment.JUSTIFY: WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def body_font_size(document: Document) -> float:
    """Tamanho do texto principal: a altura das palavras, em pontos, fica perto do corpo da fonte
    (medido no conjunto de teste: 11,2 pt de palavra para fonte de 11 pt)."""
    measured = [
        b.font_pt
        for page in document.pages
        for b in page.blocks
        if b.type is BlockType.PARAGRAPH and b.font_pt
    ]
    if measured:
        return round(min(max(median(measured), MIN_SIZE), MAX_SIZE) * 2) / 2
    heights = [
        w.bbox.h * page.page_size.w_pt / page.width / block.rel_size
        for page in document.pages
        for block in page.blocks
        if block.type is BlockType.PARAGRAPH
        for w in block.words
        if w.text != ILLEGIBLE
    ]
    if not heights:
        return DEFAULT_SIZE
    return round(min(max(median(heights), MIN_SIZE), MAX_SIZE) * 2) / 2


def _edge_align(page: Page, block: Block) -> Alignment:
    """Cabeçalho/rodapé curto (número da página) alinha pela posição dele na página."""
    if block.bbox.w >= SHORT_EDGE * page.width:
        return block.align
    center = (block.bbox.x + block.bbox.w / 2) / page.width
    if center > 0.6:
        return Alignment.RIGHT
    return Alignment.LEFT if center < 0.4 else Alignment.CENTER


class _Writer:
    def __init__(self, document: Document, assets: Path, scale: float) -> None:
        self.document = document
        self.assets = assets
        self.scale = scale
        self.base = body_font_size(document) * scale
        self.doc: Any = docx.Document()
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(self.base)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0)

    # ------------------------------------------------------------- texto

    def size(self, block: Block) -> float:
        """Tamanho medido no original (D18); sem medida, corpo vezes tamanho relativo."""
        size = block.font_pt * self.scale if block.font_pt else self.base * block.rel_size
        return round(min(max(size, MIN_SIZE), MAX_SIZE) * 2) / 2

    def runs(self, paragraph: Any, words: Sequence[Word], size: float) -> None:
        """Palavras seguidas com a mesma ênfase viram um só trecho (run)."""
        groups: list[tuple[bool, bool, list[str]]] = []
        for word in words:
            if not word.text:
                continue  # palavra removida na conferência
            style = (word.bold, word.italic)
            if groups and groups[-1][:2] == style:
                groups[-1][2].append(word.text)
            else:
                groups.append((word.bold, word.italic, [word.text]))
        for index, (bold, italic, texts) in enumerate(groups):
            text = " ".join(texts) + (" " if index < len(groups) - 1 else "")
            run = paragraph.add_run(text)
            run.bold, run.italic = bold, italic
            run.font.size = Pt(size)

    def text_paragraphs(self, container: Any, page: Page, block: Block, space_before: float) -> Any:
        size = self.size(block)
        if block.type is BlockType.LIST:
            starts = [item.word_start for item in block.items] + [len(block.words)]
            first = None
            for index, item in enumerate(block.items):
                p = container.add_paragraph()
                indent = LIST_INDENT * (item.level + 1)
                p.paragraph_format.left_indent = Pt(indent)
                p.paragraph_format.first_line_indent = Pt(-LIST_INDENT)
                self.runs(p, block.words[starts[index] : starts[index + 1]], size)
                first = first or p
            if first is not None:
                first.paragraph_format.space_before = Pt(space_before)
            return first
        p = container.add_paragraph()
        p.alignment = _ALIGN[block.align]
        p.paragraph_format.space_before = Pt(space_before)
        self.runs(p, block.words, size)
        return p

    # ------------------------------------------------------------- tabela e figura

    def table(self, page: Page, block: Block, usable_width: float, space_before: float) -> None:
        rows = max(c.row + c.row_span for c in block.cells)
        cols = max(c.col + c.col_span for c in block.cells)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(space_before)
        table = self.doc.add_table(rows=rows, cols=cols)
        table.style = "Table Grid"
        widths = self.column_widths(page, block, cols, usable_width)
        if widths:
            table.autofit = False
            # O LibreOffice segue a grade da tabela (gridCol); o Word, a largura das células.
            for index, column in enumerate(table.columns):
                column.width = Pt(widths[index])
            for row in table.rows:
                for index, cell in enumerate(row.cells):
                    cell.width = Pt(widths[index])
        for cell in block.cells:
            target = table.cell(cell.row, cell.col)
            last = table.cell(cell.row + cell.row_span - 1, cell.col + cell.col_span - 1)
            if last is not target:
                target = target.merge(last)
            self.runs(target.paragraphs[0], cell.words, self.base)

    @staticmethod
    def column_widths(page: Page, block: Block, cols: int, usable: float) -> list[float] | None:
        """Largura de cada coluna (pt) pelas células de uma coluna só, cabendo na folha."""
        found: dict[int, list[float]] = {}
        for cell in block.cells:
            if cell.col_span == 1 and cell.bbox is not None:
                found.setdefault(cell.col, []).append(cell.bbox.w)
        if len(found) != cols:
            return None
        widths = [median(found[c]) * page.page_size.w_pt / page.width for c in range(cols)]
        shrink = min(1.0, usable / sum(widths))
        return [w * shrink for w in widths]

    def picture(
        self,
        page: Page,
        block: Block,
        usable_width: float,
        space_before: float,
        lane_left: float,
    ) -> None:
        """Imagem do original na mesma posição horizontal: recuo = distância do início do
        conteúdo (`lane_left`, em px) na página original."""
        if block.image_ref is None:
            return
        k = page.page_size.w_pt / page.width
        width = min(block.bbox.w * k, usable_width)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Pt(
            min(max(0.0, (block.bbox.x - lane_left) * k), usable_width - width)
        )
        p.paragraph_format.space_before = Pt(space_before)
        p.add_run().add_picture(str(self.assets / block.image_ref), width=Pt(width))

    # ------------------------------------------------------------- páginas

    def section_for(self, page: Page) -> Any:
        if page.index == 0:
            section = self.doc.sections[0]
        else:
            section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        width, height = page.page_size.w_pt, page.page_size.h_pt
        section.orientation = WD_ORIENT.LANDSCAPE if width > height else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Pt(width), Pt(height)
        k = width / page.width
        body = [b for b in page.blocks if b.type not in {BlockType.HEADER, BlockType.FOOTER}]
        if body:
            left = min(b.bbox.x for b in body) * k
            right = width - max(b.bbox.x + b.bbox.w for b in body) * k
            top = min(b.bbox.y for b in body) * k
            bottom = height - max(b.bbox.y + b.bbox.h for b in body) * k
        else:
            left = right = top = bottom = 72.0
        margins = [min(max(m, MIN_MARGIN), MAX_MARGIN) for m in (left, right, top, bottom)]
        section.left_margin, section.right_margin = Pt(margins[0]), Pt(margins[1])
        section.top_margin, section.bottom_margin = Pt(margins[2]), Pt(margins[3])
        section.header_distance = section.footer_distance = Pt(min(margins[2], margins[3]) / 2)
        self.set_columns(section, 1)
        for part, kind in ((section.header, BlockType.HEADER), (section.footer, BlockType.FOOTER)):
            part.is_linked_to_previous = False
            for p in list(part.paragraphs):
                p.text = ""
            blocks = [b for b in page.blocks if b.type is kind]
            target = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
            for index, block in enumerate(blocks):
                p = target if index == 0 else part.add_paragraph()
                p.alignment = _ALIGN[_edge_align(page, block)]
                self.runs(p, block.words, self.size(block))
        return section

    @staticmethod
    def set_columns(section: Any, count: int) -> None:
        sect = section._sectPr
        cols = sect.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sect.append(cols)
        cols.set(qn("w:num"), str(count))
        cols.set(qn("w:space"), "425")

    def write_page(self, page: Page) -> None:
        section = self.section_for(page)
        k = page.page_size.w_pt / page.width
        usable = (section.page_width - section.left_margin - section.right_margin) / 12700
        body = sorted(
            (b for b in page.blocks if b.type not in {BlockType.HEADER, BlockType.FOOTER}),
            key=lambda b: b.order,
        )
        if not body:
            self.doc.add_paragraph()
            return
        left = min(b.bbox.x for b in body)
        content_width = max(b.bbox.x + b.bbox.w for b in body) - left
        current_columns = 1
        last_bottom: dict[int, float] = {}
        previous_column = -1
        for block in body:
            spans_all = page.columns == 1 or block.bbox.w >= FULL_WIDTH * content_width
            wanted = 1 if spans_all else page.columns
            column = 0 if spans_all else block.column
            if wanted != current_columns:
                section = self.doc.add_section(WD_SECTION.CONTINUOUS)
                self.set_columns(section, wanted)
                current_columns = wanted
                last_bottom = {}
                previous_column = -1
            # Espaço antes = vão do original, menos a folga que a própria linha já ocupa.
            above = last_bottom.get(column)
            space = 0.0 if above is None else (block.bbox.y - above) * k - LEADING_GAP * self.base
            # Vão grande (assinatura no pé da página) é mantido; a conferência de paginação
            # reduz as fontes se algo transbordar.
            space = min(max(space, 0.0), MAX_GAP * page.page_size.h_pt)
            if current_columns > 1 and previous_column >= 0 and column > previous_column:
                breaker = self.doc.add_paragraph()
                breaker.add_run().add_break(WD_BREAK.COLUMN)
            if block.type in IMAGE_TYPES:
                lane_left = (
                    left
                    if spans_all
                    else min(
                        (
                            b.bbox.x
                            for b in body
                            if b.column == column and b.bbox.w < FULL_WIDTH * content_width
                        ),
                        default=left,
                    )
                )
                self.picture(page, block, usable / max(current_columns, 1), space, lane_left)
            elif block.type is BlockType.TABLE and block.cells:
                self.table(page, block, usable / max(current_columns, 1), space)
            elif block.words:
                self.text_paragraphs(self.doc, page, block, space)
            last_bottom[column] = block.bbox.y + block.bbox.h
            previous_column = column

    def save(self, out: Path) -> Path:
        for page in self.document.pages:
            self.write_page(page)
        self.doc.save(str(out))
        return out


def build_docx(document: Document, assets: Path, out: Path, scale: float = 1.0) -> Path:
    """Gera o Word em `out`; `assets` é a pasta da conversão (onde estão `crops/` e `pages/`).
    `scale` reduz as fontes quando o PDF de conferência mostra página transbordando (D10)."""
    return _Writer(document, assets, scale).save(out)
